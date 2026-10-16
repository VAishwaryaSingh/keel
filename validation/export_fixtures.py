"""Exports every Outputs_Web value, per scenario, from the independent pycel recalculation.
The website's tests compare HyperFormula (the browser engine) with these fixtures cell by cell.
Usage: python export_fixtures.py model/keel_saas_model.xlsx web/tests/fixtures
"""
import json
import os
import sys

import openpyxl
from openpyxl.utils import get_column_letter as L
from pycel import ExcelCompiler
from pycel.excelwrapper import ExcelOpxWrapper

path, outdir = sys.argv[1:3]
os.makedirs(outdir, exist_ok=True)


def _dn(self):
    if self._defined_names is None:
        self._defined_names = {}
        for name, dn in self.workbook.defined_names.items():
            dest = [(a, sh) for sh, a in dn.destinations if sh in self.workbook]
            if dest:
                self._defined_names[str(name)] = dest
    return self._defined_names


ExcelOpxWrapper.defined_names = property(_dn)
wb = openpyxl.load_workbook(path)
ws = wb["Outputs_Web"]
xl = ExcelCompiler(filename=path)
sel = "Inputs_Drivers!" + wb.defined_names["inp_scenario"].attr_text.split("!")[1].replace("$", "")
xl.evaluate(sel)
for scen in sys.argv[3:] or ["Base", "Upside", "Downside"]:
    xl.set_value(sel, scen)
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                v = xl.evaluate(f"Outputs_Web!{c.coordinate}")
                out[c.coordinate] = v if isinstance(v, (int, float, str)) or v is None else str(v)
    json.dump(out, open(os.path.join(outdir, f"outputs_{scen.lower()}.json"), "w"))
    print(scen, len(out), "cells")
