"""Generates the hand-verification and portability kit in tests/ from the current workbook.

Usage: python make_test_kit.py model/keel_saas_model.xlsx model/snapshots.json tests/
Writes: manual_rebuild_template.xlsx, expected_values.csv, manual_verification.md, portability_checklist.md
The template contains ONLY the Base-scenario inputs and the target numbers. The calculations are for the model owner to build by hand.
"""
import csv
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill
from pycel import ExcelCompiler
from pycel.excelwrapper import ExcelOpxWrapper

path, snap_path, outdir = sys.argv[1:4]
os.makedirs(outdir, exist_ok=True)


def _dn(self):
    if self._defined_names is None:
        self._defined_names = {}
        for name, dn in self.workbook.defined_names.items():
            dest = [(a, sh) for sh, a in dn.destinations if sh in self.workbook]
            if dest:
                self._defined_names[str(name)] = dest
    return self._defined_names


ExcelOpxWrapper.defined_names = property(_dn)
wb = openpyxl.load_workbook(path)
xl = ExcelCompiler(filename=path)
snap = json.load(open(snap_path))
dws, sws = wb["Inputs_Drivers"], wb["Inputs_Scenarios"]


def find(sheet, label, nth=0):
    ws = wb[sheet]
    return [r for r in range(1, ws.max_row + 1) if ws[f"A{r}"].value == label][nth]


def ev(sheet, label, col, nth=0):
    return xl.evaluate(f"{sheet}!{col}{find(sheet, label, nth)}")


def base_input(name):
    ref = wb.defined_names[name].attr_text.split("!")[1].replace("$", "")
    r = int(re.sub(r"\D", "", ref))
    v = dws[f"C{r}"].value
    if isinstance(v, str) and v.startswith("=INDEX"):
        srow = int(re.search(r"[$]C(\d+):", v).group(1))
        v = sws["C" + str(srow)].value
    return v, dws[f"A{r}"].value, dws[f"B{r}"].value


NEEDED = ["inp_arr_open_ent", "inp_arr_open_mid", "inp_logos_ent", "inp_logos_mid", "inp_logo_growth", "inp_acv_ent", "inp_acv_mid", "inp_price_uplift",
          "inp_exp_ent", "inp_exp_mid", "inp_churn_ent", "inp_churn_mid",
          "inp_v_arr_open", "inp_v_cogs_pct", "inp_v_rd_pct", "inp_v_sm_pct", "inp_v_ga_pct", "inp_v_ap_months", "inp_v_open_cash", "inp_v_open_ar", "inp_v_open_ppe",
          "inp_ma_offer_mult", "inp_ma_cr_pct", "inp_ma_tech_pct", "inp_ma_dr_haircut", "inp_tax"]

# targets (Base scenario, computed by the workbook)
T = []
for seg, nth, nm in (("ent", 0, "Enterprise"), ("mid", 1, "Mid-market")):
    T.append((f"ARR bridge, FY1: new-logo ARR, {nm}", ev("Rev_Build", "+ New logos", "AN", nth)))
    T.append((f"ARR bridge, FY1: expansion, {nm}", ev("Rev_Build", "+ Expansion", "AN", nth)))
    T.append((f"ARR bridge, FY1: churn, {nm}", ev("Rev_Build", "- Churn", "AN", nth)))
    T.append((f"ARR bridge: closing ARR at month 12, {nm}", ev("Rev_Build", "Closing ARR", "Z", nth)))
    T.append((f"ARR bridge: closing ARR at month 24, {nm}", ev("Rev_Build", "Closing ARR", "AL", nth)))
T.append(("ARR bridge: total closing ARR at month 24", xl.evaluate(f"SaaS_Metrics!AL{find('SaaS_Metrics', 'Closing ARR')}")))
T.append(("Deferred revenue: subscription invoices raised in FY1 (both segments)", ev("Rev_Rec", "Total subscription invoices", "AN")))
T.append(("Deferred revenue: subscription revenue recognised in FY1", ev("Rev_Rec", "Total subscription revenue", "AN")))
T.append(("Deferred revenue: opening subscription deferred revenue", ev("Rev_Rec", "Subscription: closing", "N")))
T.append(("Deferred revenue: closing subscription deferred revenue, month 12", ev("Rev_Rec", "Subscription: closing", "Z")))
T.append(("Deferred revenue: closing subscription deferred revenue, month 24", ev("Rev_Rec", "Subscription: closing", "AL")))
T.append(("Goodwill: equity purchase price", xl.evaluate(f"MA_PPA!C{find('MA_PPA', 'Equity purchase price')}")))
T.append(("Goodwill: fair value of identifiable net assets, IFRS 3", xl.evaluate(f"MA_PPA!C{find('MA_PPA', 'Fair value of identifiable net assets')}")))
T.append(("Goodwill: IFRS 3", xl.evaluate(f"MA_PPA!C{find('MA_PPA', 'Goodwill = equity purchase price - fair value of net assets')}")))
T.append(("Goodwill: US GAAP (ASU 2021-08)", xl.evaluate(f"MA_PPA!D{find('MA_PPA', 'Goodwill = equity purchase price - fair value of net assets')}")))

# ---- template workbook
t = openpyxl.Workbook()
bold = Font(name="Arial", size=10, bold=True)
norm = Font(name="Arial", size=10)
blue = Font(name="Arial", size=10, color="0000FF")
fill_in = PatternFill("solid", fgColor="FFF9E0")
hdr_fill = PatternFill("solid", fgColor="1F3A5F")
wsI = t.active
wsI.title = "Inputs"
wsI["A1"], wsI["A1"].font = "Base-scenario inputs (typed values, copied from the model on the day this kit was made)", bold
wsI["A2"], wsI["A2"].font = "Build your calculations on new sheets that point at these cells. Do not look at the model's formulas first.", Font(name="Arial", size=10, italic=True)
for j, h in enumerate(["Name in the model", "Label", "Unit", "Value"], start=1):
    c = wsI.cell(row=4, column=j, value=h)
    c.font, c.fill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), hdr_fill
for i, nm in enumerate(NEEDED):
    v, lab, unit = base_input(nm)
    r = 5 + i
    wsI.cell(row=r, column=1, value=nm).font = norm
    wsI.cell(row=r, column=2, value=lab).font = norm
    wsI.cell(row=r, column=3, value=unit).font = norm
    c = wsI.cell(row=r, column=4, value=v)
    c.font, c.fill = blue, fill_in
wsI.column_dimensions["A"].width, wsI.column_dimensions["B"].width, wsI.column_dimensions["C"].width = 24, 54, 22
wsC = t.create_sheet("Compare")
wsC["A1"], wsC["A1"].font = "Compare your rebuild with the model", bold
wsC["A2"], wsC["A2"].font = "Type, in column C, a formula that points at your own calculation. Columns D and E fill in by themselves.", Font(name="Arial", size=10, italic=True)
for j, h in enumerate(["Item", "Model value (GBP k)", "Your value", "Difference", "Status"], start=1):
    c = wsC.cell(row=4, column=j, value=h)
    c.font, c.fill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), hdr_fill
for i, (lab, v) in enumerate(T):
    r = 5 + i
    wsC.cell(row=r, column=1, value=lab).font = norm
    c = wsC.cell(row=r, column=2, value=round(v, 6))
    c.font, c.number_format = norm, '#,##0.00;(#,##0.00)'
    wsC.cell(row=r, column=3).fill = fill_in
    wsC.cell(row=r, column=3).number_format = '#,##0.00;(#,##0.00)'
    wsC.cell(row=r, column=4, value=f'=IF(C{r}="","",C{r}-B{r})').number_format = '#,##0.0000;(#,##0.0000)'
    wsC.cell(row=r, column=5, value=f'=IF(C{r}="","not yet",IF(ABS(C{r}-B{r})<0.01,"MATCH","DIFFERENT"))')
last = 4 + len(T)
wsC.cell(row=last + 2, column=1, value="Items matched").font = bold
wsC.cell(row=last + 2, column=2, value=f'=COUNTIF(E5:E{last},"MATCH")&" of {len(T)}"').font = bold
wsC.column_dimensions["A"].width, wsC.column_dimensions["B"].width, wsC.column_dimensions["C"].width = 70, 20, 16
wsC.column_dimensions["D"].width, wsC.column_dimensions["E"].width = 14, 14
for nm_ in ("Your_ARR", "Your_DeferredRevenue", "Your_Goodwill"):
    w = t.create_sheet(nm_)
    w["A1"], w["A1"].font = f"{nm_.replace('_', ' ')}: build it here", bold
t.save(os.path.join(outdir, "manual_rebuild_template.xlsx"))

# ---- expected values csv
mets = [(k, snap["scen"]["Base"][k]) for k in snap["scen"]["Base"]]
labels = {}
sp = json.load(open(os.path.join(os.path.dirname(snap_path), "snap_spec.json")))
with open(os.path.join(outdir, "expected_values.csv"), "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["metric_id", "Base", "Upside", "Downside", "Custom (= Base)"])
    for k in snap["scen"]["Base"]:
        row = [k] + [snap["scen"][s][k] for s in ("Base", "Upside", "Downside")] + [snap["scen"]["Base"][k]]
        w.writerow([round(x, 4) if isinstance(x, float) else x for x in row])

# ---- documents
rows_md = "\n".join(f"| {lab} | {v:,.2f} |" for lab, v in T)
open(os.path.join(outdir, "manual_verification.md"), "w").write(f"""# Manual verification (build principle 10; Phase 5 task)

**Goal:** rebuild three calculations from the inputs alone, on blank sheets, and confirm they match the model. The point is that you understand them, not that a script agrees. Do **not** read the workbook's formulas first, and do not ask a tool to write the formulas for you.

**Kit:** `tests/manual_rebuild_template.xlsx` holds the Base-scenario inputs (sheet `Inputs`), the model's answers (sheet `Compare`, column B), and three empty sheets for your work. Type into `Compare` column C a formula that points at your result; columns D and E report the difference and `MATCH` / `DIFFERENT`. Tolerance: 0.01 GBP k.

## 1. ARR bridge, months 1 to 24, by segment

For each segment, month by month:

- *Volume.* New logos per month grow at the annual new-logo growth rate, compounding monthly, so month t has `base logos x (1 + growth)^((t-1)/12)`.
- *Price.* The ACV of a new logo steps up by the price uplift at the start of each model year (months 13-24 have one uplift).
- *New-logo ARR* = logos x ACV.
- *Rates.* The annual churn and expansion inputs are converted to monthly compounding rates: monthly churn = 1 - (1 - annual)^(1/12); monthly expansion = (1 + annual)^(1/12) - 1.
- *Expansion* and *churn* in a month are the monthly rates applied to that month's **opening** ARR (churn is a negative).
- Closing ARR = opening + new + expansion + churn. Opening ARR of month 1 is the opening input.

## 2. Subscription deferred revenue roll-forward

Both segments are billed **annually upfront** and recognised **evenly over 12 months, starting in the month of invoice**.

- *Invoices.* In month t the invoice is that month's new-logo ARR (annual price paid upfront) **plus** the renewal of the invoice raised 12 months earlier, scaled by (1 + annual expansion - annual churn). For months 1 to 12, the invoice "12 months earlier" is the opening ARR divided by 12 (opening-year customers renew evenly through the year).
- *Revenue* in month t = one twelfth of the sum of the invoices raised in months t-11 to t (including months before the model starts: each of those 12 history months is opening ARR / 12).
- *Opening deferred revenue.* An invoice raised k months before the opening date (k = 0 for the last month) still has (11 - k)/12 of its value unrecognised at the opening date. Sum that across the 12 history invoices.
- Roll-forward: closing = opening + invoices - revenue.

## 3. Goodwill (IFRS 3 and US GAAP)

Use the Vaultline inputs on the `Inputs` sheet. Steps in words:

- Vaultline's opening deferred revenue is built the same way as in section 2 (opening ARR / 12 per history month, same unexpired fractions).
- Payables = opening ARR / 12 x (sum of the four cost percentages) x payables months.
- Book net assets = cash + receivables + PP&E - payables - deferred revenue.
- Enterprise value = offer multiple x opening ARR. Equity price = EV + cash acquired.
- Customer relationships and technology are percentages of EV. A deferred tax liability = tax rate x those two intangibles.
- IFRS 3 writes deferred revenue down by the haircut percentage (the liability falls, so net assets rise). US GAAP after ASU 2021-08 does not.
- Goodwill = equity price - fair value of identifiable net assets.

## Targets (Base scenario, GBP k)

| Item | Model value |
|---|---|
{rows_md}

## When something does not match

Find the first month or line where you and the model diverge, then work out which of you is wrong. Either answer is useful: a bug in the model belongs in `docs/ai_usage.md`; a gap in your understanding belongs in your own notes before you defend this model to anyone. Record the outcome below.

| Date | Section | Matched? | Notes / what you found |
|---|---|---|---|
| | ARR bridge | | |
| | Deferred revenue | | |
| | Goodwill | | |
""")

_names = json.load(open(os.path.join(os.path.dirname(snap_path), "named_ranges.json")))
NNAMES, NINP, NOUT = len(_names), sum(k.startswith("inp_") for k in _names), sum(k.startswith("out_") for k in _names)
exp = {s: snap["scen"][s] for s in ("Base", "Upside", "Downside")}
pick = [("rev_fy1", "FY1 revenue"), ("ebitda_fy1", "FY1 EBITDA"), ("arr24", "ARR, month 24"), ("nrr12", "NRR, month 12"), ("cash24", "Cash, month 24"),
        ("mincash", "Lowest month-end cash"), ("ma_value", "Deal value created"), ("pf_cash24", "Pro-forma cash, month 24")]
tbl = "\n".join("| " + lab + " | " + " | ".join((f"{exp[s][k]:.4f}" if abs(exp[s][k]) < 5 else f"{exp[s][k]:,.0f}") for s in ("Base", "Upside", "Downside")) + " |" for k, lab in pick)
open(os.path.join(outdir, "portability_checklist.md"), "w").write(f"""# Portability checklist (Excel and Google Sheets)

The workbook was recalculated in Python (`pycel`), never in Excel or Google Sheets. This checklist is how to close that gap. Tick each box and note any difference. Full expected values for every metric and scenario are in `tests/expected_values.csv`.

## Excel

- [ ] Open `model/keel_saas_model.xlsx` and let it calculate (it is set to recalculate on open). No `#NAME?`, `#REF!`, `#VALUE!` or `#DIV/0!` anywhere (Ctrl+F for `#`).
- [ ] `Cover!B4` reads **ALL CHECKS PASS** and `Cover!B7` reads 0 outlier warnings.
- [ ] Formulas > Name Manager shows {NNAMES} names ({NINP} `inp_*`, {NOUT} `out_*`, plus `chk_*` and `scen_*`).
- [ ] Change the dropdown on `Inputs_Drivers!C4` to each scenario and compare with the table below (Downside: master check reads **CHECK FAILURE** by design, because pro-forma cash goes negative).
- [ ] Type 0.9 into the tax-rate input: Excel should refuse it (data validation, allowed range 0 to 0.5).
- [ ] `MA_Valuation`: the football field chart renders, with floating bars.
- [ ] Columns C:M are grouped and collapsed on the time-series tabs; the outline buttons expand them. `Rev_Rec` shows them expanded.
- [ ] Switch to **Custom**, change a value in `Inputs_Scenarios` column F, and confirm the model responds.
- [ ] Save as .xlsx and re-open: nothing changes.

## Google Sheets

- [ ] File > Import > Upload > "Replace spreadsheet". Wait for calculation. No error values.
- [ ] Same Cover, scenario table, and data-validation checks as above.
- [ ] Data > Named ranges lists the names; `Outputs_Web` links resolve.
- [ ] The football field chart renders (Sheets may restyle it).
- [ ] Note any function Sheets treats differently. The workbook uses only 21 classic functions (ABS, AND, AVERAGE, COLUMN, COUNT, COUNTIF, EOMONTH, IF, INDEX, ISNUMBER, LARGE, MATCH, MAX, MIN, OR, ROUNDUP, ROW, SUM, SUMIF, SUMPRODUCT, YEAR).

## Expected values (from the independent recalculation; GBP k unless a ratio)

| Metric | Base | Upside | Downside |
|---|---|---|---|
{tbl}

## Result

| Engine | Date | Matches? | Differences found |
|---|---|---|---|
| Excel | | | |
| Google Sheets | | | |
| HyperFormula (Phase 6) | | | |
""")
print("kit written to", outdir, f"({len(T)} targets)")
