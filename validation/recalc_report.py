"""Independent recalculation of the workbook with `pycel` (no Excel needed).

Usage: python recalc_report.py <model.xlsx> [Base|Upside|Downside|Custom]
Sets the scenario selector on a temporary copy, evaluates the cells lazily, prints checks and key outputs.
Dev tool: requires `pip install openpyxl pycel` in a virtualenv.
"""
import os
import sys
import tempfile
import openpyxl
from pycel import ExcelCompiler
from pycel.excelwrapper import ExcelOpxWrapper


def _defined_names(self):
    # pycel predates the openpyxl 3.1 defined-names API; rebuild the mapping it expects.
    if self._defined_names is None:
        self._defined_names = {}
        for name, dn in self.workbook.defined_names.items():
            dest = [(alias, sheet) for sheet, alias in dn.destinations if sheet in self.workbook]
            if dest:
                self._defined_names[str(name)] = dest
    return self._defined_names


ExcelOpxWrapper.defined_names = property(_defined_names)

path = sys.argv[1]
scen = sys.argv[2] if len(sys.argv) > 2 else "Base"

wb = openpyxl.load_workbook(path)
sel = wb.defined_names["inp_scenario"].attr_text.split("!")[1].replace("$", "")
wb["Inputs_Drivers"][sel] = scen
tmp = os.path.join(tempfile.mkdtemp(), "m.xlsx")
wb.save(tmp)
xl = ExcelCompiler(filename=tmp)


def val(sheet, ref):
    return xl.evaluate(f"{sheet}!{ref}")


def find(sheet, label):
    ws = wb[sheet]
    for r in range(1, ws.max_row + 1):
        if ws[f"A{r}"].value == label:
            return r
    raise KeyError(label)


print(f"=== Scenario: {scen} ===")
ws = wb["Checks"]
for r in range(5, ws.max_row + 1):
    if ws[f"A{r}"].value == "MASTER CHECK":
        print(f"{'MASTER CHECK':34s} {val('Checks', f'E{r}')}")
        continue
    if ws[f"A{r}"].value == "WARNINGS":
        print(f"{'WARNINGS (series flagged)':34s} {val('Checks', f'E{r}')}")
        continue
    if isinstance(ws[f"C{r}"].value, str) and ws[f"C{r}"].value.startswith("=") and ws[f"E{r}"].value and "IF(ABS" in str(ws[f"E{r}"].value):
        print(f"{ws[f'A{r}'].value:34s} value={val('Checks', f'C{r}'):>12.6f}  pass={val('Checks', f'E{r}')}")
    elif isinstance(ws[f"E{r}"].value, str) and "REVIEW" in ws[f"E{r}"].value:
        print(f"  warning: {ws[f'A{r}'].value:40s} months flagged={val('Checks', f'C{r}')}  {val('Checks', f'E{r}')}")

ROWS = (("IS", "Total revenue", "{:,.0f}"), ("IS", "EBITDA", "{:,.0f}"), ("IS", "EBITDA margin", "{:.1%}"),
        ("IS", "Gross margin", "{:.1%}"), ("IS", "Net income", "{:,.0f}"),
        ("BS", "Cash", "{:,.0f}"), ("BS", "Deferred revenue", "{:,.0f}"),
        ("SaaS_Metrics", "Closing ARR", "{:,.0f}"), ("SaaS_Metrics", "ARR growth, year on year", "{:.1%}"),
        ("SaaS_Metrics", "Net revenue retention (NRR)", "{:.1%}"), ("SaaS_Metrics", "Gross revenue retention (GRR)", "{:.1%}"),
        ("SaaS_Metrics", "CAC payback (months of gross-margin on new + expansion ARR)", "{:.1f}"),
        ("SaaS_Metrics", "LTV / CAC", "{:.1f}"),
        ("SaaS_Metrics", "Burn multiple (net burn / net new ARR)", "{:.2f}"),
        ("SaaS_Metrics", "Cash runway (months at trailing 3m average burn)", "{:.1f}"),
        ("SaaS_Metrics", "Rule of 40 (ARR growth + EBITDA margin)", "{:.1%}"),
        ("Headcount_Opex", "Total closing FTE", "{:,.0f}"),
        ("Rev_Rec", "Subscription revenue / (closing ARR / 12)", "{:.1%}"))
for sheet, label, fmt in ROWS:
    r = find(sheet, label)
    out = []
    for col in ("O", "Z", "AL"):          # month 1, 12, 24
        x = val(sheet, f"{col}{r}")
        try:
            out.append(fmt.format(float(x)))
        except (TypeError, ValueError):
            out.append(str(x))
    print(f"{sheet + ': ' + label[:40]:50s} M1={out[0]:>10} M12={out[1]:>10} M24={out[2]:>10}")
